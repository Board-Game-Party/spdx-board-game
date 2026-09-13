import io
import csv
from datetime import datetime, timezone
from typing import Optional
from sqlalchemy.orm import Session
from fastapi import HTTPException, status
from fastapi.responses import Response

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill

from backend.app.shared.models import (
    Assignment, Criterion, ClassroomMember, User, GroupEntity,
    PairAssignment, Comparison, AuditEvent
)
from backend.app.features.reporting.services import (
    get_group_summary_report, get_individual_summary_report, get_pair_coverage_report
)

def export_assignment_csv(
    db: Session,
    assignment: Assignment,
    report_type: str,  # group | individual | coverage | raw
    mask_identities: bool = True,
    current_user: Optional[User] = None
) -> Response:
    """
    US-EXPORT-01 / FR-EXPORT-01 / FR-EXPORT-03..05:
    Export to UTF-8 BOM CSV for Excel Thai support.
    """
    output = io.StringIO()
    # Write UTF-8 BOM
    output.write("\ufeff")
    writer = csv.writer(output)

    now_str = datetime.now().strftime("%Y%m%d-%H%M")
    filename = f"{assignment.classroom.slug}_{assignment.slug}_{report_type}_{now_str}.csv"

    if report_type == "group":
        rep = get_group_summary_report(db, assignment)
        writer.writerow(["Group ID", "Group Name", "Component Score", "Max Score", "Comparisons", "Flags", "Overridden", "Override Score", "Override Reason"])
        for g in rep.groups:
            writer.writerow([
                g.group_id, g.group_name, round(g.component_score, 3), g.max_score,
                g.comparison_count, ";".join(g.flags), g.is_overridden, g.override_score or "", g.override_reason or ""
            ])

    elif report_type == "individual":
        rep = get_individual_summary_report(db, assignment)
        writer.writerow([
            "Student ID", "Display Name", "Email", "Group Name",
            "Group Score", "Individual Score", "Participation Ratio (p)", "Multiplier (M)",
            "Net Final Score", "Max Score", "Flags", "Overridden", "Override Reason"
        ])
        for s in rep.students:
            writer.writerow([
                s.student_id or "", s.display_name, s.email, s.group_name or "",
                round(s.group_component_score, 3), round(s.individual_component_score, 3),
                round(s.participation_ratio, 3), round(s.participation_multiplier, 3),
                round(s.net_final_score, 3), s.max_possible_score, ";".join(s.flags),
                s.is_overridden, s.override_score or "", s.override_reason or ""
            ])

    elif report_type == "coverage":
        rep = get_pair_coverage_report(db, assignment)
        writer.writerow(["Criterion", "Side", "Item A", "Item B", "Target Coverage", "Actual Coverage", "Mean Choice", "Low Coverage Flag"])
        for p in rep.pairs:
            writer.writerow([
                p.criterion_name, p.side, p.item_a_name, p.item_b_name,
                p.target_coverage, p.actual_coverage, p.mean_choice or "", p.is_low_coverage
            ])

    elif report_type == "raw":
        # FR-EXPORT-03 / FR-EXPORT-04: Mask identities by default
        pairs = db.query(PairAssignment).filter(PairAssignment.assignment_id == assignment.id).all()
        evaluator_map = {}
        eval_counter = 1

        writer.writerow([
            "Pair Assignment ID", "Criterion ID", "Side", "Item A ID", "Item B ID",
            "Display Left Item ID", "Evaluator ID", "Choice (1-6)", "Status", "Time on Task (ms)"
        ])
        for p in pairs:
            comp = db.query(Comparison).filter(Comparison.pair_assignment_id == p.id).first()
            raw_eval_id = p.evaluator_user_id
            if mask_identities:
                if raw_eval_id not in evaluator_map:
                    evaluator_map[raw_eval_id] = f"evaluator_{eval_counter:03d}"
                    eval_counter += 1
                eval_id_out = evaluator_map[raw_eval_id]
            else:
                eval_id_out = raw_eval_id

            writer.writerow([
                p.id, p.criterion_id, p.side, p.item_a_id, p.item_b_id,
                p.display_left_item_id, eval_id_out,
                comp.choice if comp else "",
                comp.status if comp else "DRAFT",
                comp.time_on_task_ms if comp else 0
            ])

        # Log audit if exporting unmasked identities (FR-EXPORT-04)
        if not mask_identities and current_user:
            audit = AuditEvent(
                classroom_id=assignment.classroom_id,
                assignment_id=assignment.id,
                actor_user_id=current_user.id,
                action="EXPORT_IDENTITIES",
                resource_type="ASSIGNMENT",
                resource_id=assignment.id,
                reason="Export raw comparisons with real evaluator identities"
            )
            db.add(audit)
            db.commit()

    csv_bytes = output.getvalue().encode("utf-8")
    return Response(
        content=csv_bytes,
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )

def export_assignment_xlsx(
    db: Session,
    assignment: Assignment,
    mask_identities: bool = True,
    current_user: Optional[User] = None
) -> Response:
    """
    US-EXPORT-01 / FR-EXPORT-02 / FR-EXPORT-06:
    Export to multi-sheet XLSX file (Group, Individual, Coverage, Metadata).
    """
    wb = openpyxl.Workbook()
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    # 1. Sheet: Group Summary
    ws_grp = wb.active
    ws_grp.title = "Group Summary"
    grp_rep = get_group_summary_report(db, assignment)
    headers_grp = ["Group ID", "Group Name", "Component Score", "Max Score", "Comparisons", "Flags", "Overridden", "Override Score", "Override Reason"]
    ws_grp.append(headers_grp)
    for g in grp_rep.groups:
        ws_grp.append([
            g.group_id, g.group_name, round(g.component_score, 3), g.max_score,
            g.comparison_count, ";".join(g.flags), g.is_overridden, g.override_score or "", g.override_reason or ""
        ])

    # 2. Sheet: Individual Summary
    ws_ind = wb.create_sheet(title="Individual Summary")
    ind_rep = get_individual_summary_report(db, assignment)
    headers_ind = [
        "Student ID", "Display Name", "Email", "Group Name",
        "Group Score", "Individual Score", "Participation Ratio (p)", "Multiplier (M)",
        "Net Final Score", "Max Score", "Flags", "Overridden", "Override Reason"
    ]
    ws_ind.append(headers_ind)
    for s in ind_rep.students:
        ws_ind.append([
            s.student_id or "", s.display_name, s.email, s.group_name or "",
            round(s.group_component_score, 3), round(s.individual_component_score, 3),
            round(s.participation_ratio, 3), round(s.participation_multiplier, 3),
            round(s.net_final_score, 3), s.max_possible_score, ";".join(s.flags),
            s.is_overridden, s.override_score or "", s.override_reason or ""
        ])

    # 3. Sheet: Pair Coverage
    ws_cov = wb.create_sheet(title="Pair Coverage")
    cov_rep = get_pair_coverage_report(db, assignment)
    headers_cov = ["Criterion", "Side", "Item A", "Item B", "Target Coverage", "Actual Coverage", "Mean Choice", "Low Coverage Flag"]
    ws_cov.append(headers_cov)
    for p in cov_rep.pairs:
        ws_cov.append([
            p.criterion_name, p.side, p.item_a_name, p.item_b_name,
            p.target_coverage, p.actual_coverage, p.mean_choice or "", p.is_low_coverage
        ])

    # 4. Sheet: Metadata (FR-EXPORT-06)
    ws_meta = wb.create_sheet(title="Metadata")
    ws_meta.append(["Property", "Value"])
    ws_meta.append(["Classroom", assignment.classroom.name])
    ws_meta.append(["Classroom Slug", assignment.classroom.slug])
    ws_meta.append(["Assignment Name", assignment.name])
    ws_meta.append(["Assignment Slug", assignment.slug])
    ws_meta.append(["Status", assignment.status])
    ws_meta.append(["Scoring Algorithm Version", assignment.scoring_formula_version])
    ws_meta.append(["Score Floor", float(assignment.score_floor)])
    ws_meta.append(["Score Ceiling", float(assignment.score_ceiling)])
    ws_meta.append(["Completion Threshold", float(assignment.completion_threshold)])
    ws_meta.append(["Instructor Weight", float(assignment.instructor_weight)])
    ws_meta.append(["Export Timestamp (UTC)", datetime.now(timezone.utc).isoformat()])

    # Style headers on all sheets
    for ws in [ws_grp, ws_ind, ws_cov, ws_meta]:
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    out_stream = io.BytesIO()
    wb.save(out_stream)
    out_stream.seek(0)

    now_str = datetime.now().strftime("%Y%m%d-%H%M")
    filename = f"{assignment.classroom.slug}_{assignment.slug}_full_report_{now_str}.xlsx"

    return Response(
        content=out_stream.read(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"}
    )
